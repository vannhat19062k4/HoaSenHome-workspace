"""Join warehouse allocations to store pins without inventing unmatched locations.

Usage: python3 build_supply.py '/path/to/Data thống kê CH 3 miền.xlsx'
"""
import collections
import difflib
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

if len(sys.argv) != 2:
    raise SystemExit(__doc__)

ROOT = Path(__file__).parent
STORES = json.loads((ROOT / 'data/stores.json').read_text())
SHEET = openpyxl.load_workbook(sys.argv[1], data_only=True).active

WAREHOUSES = [
    ('binh-duong', 'TK Bình Dương', 'Nam', 'Thống Nhất/Số 9 2, KCN Sóng Thần, Dĩ An, Hồ Chí Minh, Việt Nam'),
    ('phan-thiet', 'TK Phan Thiết', 'Nam', 'Lô 4/1 Khu Công Nghiệp Phan Thiết Giai Đoạn 1, Xã Phong Nẫm, Thành phố Phan Thiết, Tỉnh Bình Thuận'),
    ('cai-cui', 'TK Cái Cui', 'Nam', 'Số 2 Khu vực Phú Thắng, Phường Hưng Phú, Thành phố Cần Thơ'),
    ('binh-dinh', 'TK Bình Định', 'Trung', 'Khu A1.2 và A1.3.2, Khu công nghiệp Nhơn Hòa, Phường An Nhơn Nam, Tỉnh Bình Định'),
    ('dak-lak', 'TK Đắk Lắk', 'Trung', 'Lô C6, Cụm Công nghiệp Tân An 1, Phường Tân An, TP. Buôn Ma Thuột, Tỉnh Đắk Lắk'),
    ('da-nang', 'TK Đà Nẵng', 'Trung', 'Đường Số 2, KCN Hòa Cầm, Cẩm Lệ, Đà Nẵng, Việt Nam'),
    ('nghe-an', 'TK Nghệ An', 'Bắc', 'Lô CN 1-8, Khu công nghiệp Đông Hồi, Xã Quỳnh Lập, Thị xã Hoàng Mai, Tỉnh Nghệ An'),
    ('ha-nam', 'TK Hà Nam', 'Bắc', 'Khu công nghiệp Thanh Liêm, phường Châu Sơn, thành phố Phủ Lý, tỉnh Hà Nam'),
    ('yen-bai', 'TK Yên Bái', 'Bắc', 'Tổ dân phố Đồng Danh, Phường Âu Lâu, Tỉnh Yên Bái'),
]
BY_NAME = {name: slug for slug, name, _, _ in WAREHOUSES}
# Reference points from the named industrial park/port or surrounding area on Maps.
# These are deliberately not presented as verified warehouse gates.
WAREHOUSE_POINTS = {
    'binh-duong': (10.9177067, 106.7437961),
    'phan-thiet': (10.9512198, 108.1014347),
    'cai-cui': (9.9835018, 105.8357661),
    'binh-dinh': (13.8527750, 109.0744000),
    'dak-lak': (12.7335665, 108.0815553),
    'da-nang': (16.0083683, 108.1849231),
    'nghe-an': (19.2729982, 105.7751696),
    'ha-nam': (20.4977133, 105.9075142),
    'yen-bai': (21.6879285, 104.8396248),
}
PROVINCE_ALIASES = {
    'Kiên Giang':'An Giang', 'Hậu Giang':'Cần Thơ', 'Sóc Trăng':'Cần Thơ',
    'Trà Vinh':'Vĩnh Long', 'Bến Tre':'Vĩnh Long', 'Long An':'Tây Ninh',
    'Bình Dương':'Hồ Chí Minh', 'Bà Rịa - Vũng Tàu':'Hồ Chí Minh',
    'Bình Phước':'Đồng Nai', 'Bình Thuận':'Lâm Đồng', 'Đắk Nông':'Lâm Đồng',
    'Ninh Thuận':'Khánh Hòa', 'Phú Yên':'Đắk Lắk', 'Bình Định':'Gia Lai',
    'Kon Tum':'Quảng Ngãi', 'Quảng Nam':'Đà Nẵng', 'Quảng Bình':'Quảng Trị',
    'Hà Nam':'Ninh Bình', 'Nam Định':'Ninh Bình', 'Thái Bình':'Hưng Yên',
    'Hải Dương':'Hải Phòng', 'Bắc Giang':'Bắc Ninh', 'Vĩnh Phúc':'Phú Thọ',
    'Hòa Bình':'Phú Thọ', 'Yên Bái':'Lào Cai', 'Hà Giang':'Tuyên Quang',
    'Bắc Kạn':'Thái Nguyên', 'TP.HCM':'Hồ Chí Minh',
}
# Worksheet row and regional block identify a record even when names repeat.
REVIEWED_MATCHES = {
    ('Nam',40):'home:N69', ('Nam',69):'home:N57',
    ('Nam',90):'home:L08', ('Nam',95):'home:C28',
    ('Nam',114):'home:D07', ('Nam',133):'traditional:D03',
    ('Nam',145):'traditional:D17', ('Nam',155):'home:N80',
    ('Nam',46):'traditional:A40', ('Nam',120):'home:L37',
    ('Nam',127):'home:N39',
    ('Trung',39):'home:E39', ('Trung',43):'home:E46',
    ('Trung',45):'home:M09', ('Trung',65):'home:F09',
    ('Trung',99):'home:M32',
    ('Bắc',13):'traditional:G33', ('Bắc',88):'traditional:TX04',
    ('Bắc',122):'traditional:TX03',
}

def clean(value):
    return ' '.join(str(value).split()) if value is not None else ''

def norm(value):
    value = unicodedata.normalize('NFD', clean(value).casefold())
    value = ''.join(c for c in value if unicodedata.category(c) != 'Mn').replace('đ','d')
    value = re.sub(r'\b(tp|thanh pho|thi xa|cua hang|sieu thi)\.?\s*','',value)
    return re.sub(r'[^a-z0-9]+',' ',value).strip()

def score(row, store):
    if row['region'] != store['region'] or row['province'] != store['province']:
        return 0.0
    expected_type = 'traditional' if row['model'] == 'CHTT' else 'home'
    if store['type'] != expected_type:
        return 0.0
    n = norm(row['name'])
    alternatives = {norm(store['name']), norm(store['groupName'])} - {''}
    if n in alternatives:
        return 1.0
    return max((difflib.SequenceMatcher(None,n,other).ratio() for other in alternatives), default=0.0)

rows = []
for region, cols in [('Nam',(2,3,4,5,6)), ('Trung',(8,9,10,11,12)), ('Bắc',(14,15,16,17,18))]:
    for rowno in range(11,SHEET.max_row+1):
        old,name,model,primary,secondary = (clean(SHEET.cell(rowno,c).value) for c in cols)
        if not (name and model and primary.startswith('TK ')):
            continue
        if primary not in BY_NAME:
            raise ValueError(f'Unknown warehouse {primary!r} at {region} row {rowno}')
        alternates = [BY_NAME[x.strip()] for x in secondary.split(',') if x.strip() in BY_NAME and x.strip() != primary]
        rows.append({'id':f'{region.lower()}-{rowno}', 'sourceRow':rowno,
                     'region':region, 'oldProvince':old, 'province':PROVINCE_ALIASES.get(old,old),
                     'name':name, 'model':model, 'primaryId':BY_NAME[primary],
                     'alternateIds':list(dict.fromkeys(alternates)), 'storeId':None, 'match':'unmatched'})

store_by_id = {s['id']:s for s in STORES}
used = set()
# Exact singletons are reliable. A reviewed mapping resolves repeated legacy names.
for row in rows:
    key = (row['region'],row['sourceRow'])
    if key in REVIEWED_MATCHES:
        sid = REVIEWED_MATCHES[key]
        if sid not in store_by_id or sid in used:
            raise ValueError(f'Invalid reviewed match {key}: {sid}')
        row['storeId'],row['match'] = sid,'reviewed'
        used.add(sid)

for row in rows:
    if row['storeId']:
        continue
    candidates = [s for s in STORES if s['id'] not in used and score(row,s)==1.0]
    if len(candidates)==1:
        row['storeId'],row['match'] = candidates[0]['id'],'exact'
        used.add(row['storeId'])

# Match minor spelling/zero-padding differences only when there is a clear winner.
for row in rows:
    if row['storeId']:
        continue
    candidates = sorted(((score(row,s),s) for s in STORES if s['id'] not in used),key=lambda x:x[0],reverse=True)
    if not candidates:
        continue
    top,store = candidates[0]
    second = candidates[1][0] if len(candidates)>1 else 0
    if top>=0.88 and top-second>=0.07:
        row['storeId'],row['match'] = store['id'],'normalized'
        used.add(row['storeId'])

unmatched=[{k:v for k,v in row.items() if k!='alternateIds'} for row in rows if not row['storeId']]
warehouses=[{'id':slug,'name':name,'region':region,'address':address,
             'lat':WAREHOUSE_POINTS[slug][0],'lng':WAREHOUSE_POINTS[slug][1],
             'mapsUrl':'','coordinateStatus':'area_reference'}
            for slug,name,region,address in WAREHOUSES]
payload={'warehouses':warehouses,'assignments':rows}
(ROOT/'data/supply.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2))
(ROOT/'data/supply.js').write_text('window.HOA_SEN_SUPPLY = '+json.dumps(payload,ensure_ascii=False,separators=(',',':'))+';\n')
audit={'assignments':len(rows),'linked':len(rows)-len(unmatched),'unmatched':unmatched,
       'warehouseCounts':dict(collections.Counter(r['primaryId'] for r in rows)),
       'regionCounts':dict(collections.Counter(r['region'] for r in rows)),
       'unassignedStoreIds':[s['id'] for s in STORES if s['id'] not in used],
       'modelConflicts':[{'sourceRow':r['sourceRow'],'region':r['region'],'name':r['name'],
                          'supplyModel':r['model'],'storeId':r['storeId'],
                          'locationType':store_by_id[r['storeId']]['type']}
                         for r in rows if r['storeId'] and (r['model']=='CHTT') != (store_by_id[r['storeId']]['type']=='traditional')],
       'note':'Warehouse pins are reference points for the named industrial park/port/area, not verified warehouse gates. Driving distance is calculated on demand in the app; trucking restrictions are not modeled.'}
(ROOT/'data/supply_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in audit.items() if k not in ('unmatched','unassignedStoreIds')},ensure_ascii=False,indent=2))
print('Unmatched:',[(x['region'],x['sourceRow'],x['name']) for x in unmatched])
